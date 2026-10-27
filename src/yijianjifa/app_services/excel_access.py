"""Frozen native Excel previews, explicit approvals and atomic durable results."""
from dataclasses import asdict
import hashlib
import json
import secrets
import threading

from sqlalchemy import text
from openpyxl import load_workbook
from yijianjifa.features.library.copy_service import CopyRepository,CategoryRepository
from yijianjifa.features.library.excel_import import read_excel_rows,classify_rows,apply_import


class ExcelAccess:
    def __init__(self,engine):
        self._engine=engine;self._saved_files={};self._previews={};self._lock=threading.RLock()

    def register_saved(self,path):
        from pathlib import Path
        target=Path(path).resolve(strict=True)
        if target.suffix.lower()!='.xlsx' or not target.is_file():raise ValueError('文件无效。')
        token=secrets.token_hex(16)
        with self._lock:
            if len(self._saved_files)>=32:self._saved_files.pop(next(iter(self._saved_files)))
            self._saved_files[token]=target
        return {'saved':True,'name':target.name,'token':token}

    def open_saved(self,token,action):
        if type(token) is not str or action not in ('file','folder'):raise ValueError('请求无效。')
        with self._lock:target=self._saved_files.get(token)
        if target is None or not target.is_file():raise ValueError('文件已移动或不存在，请重新下载。')
        if action=='folder':
            from yijianjifa.os_integration.reveal import reveal_path
            return {'opened':reveal_path(target)}
        import platform,subprocess,os
        system=platform.system()
        if system=='Windows':os.startfile(str(target))
        else:subprocess.Popen(['open' if system=='Darwin' else 'xdg-open',str(target)])
        return {'opened':True}

    def preview(self,path):
        workbook=load_workbook(path,read_only=True,data_only=False)
        try:
            for row in workbook.worksheets[0].iter_rows():
                if any(cell.data_type=='f' for cell in row):
                    raise ValueError('表格含公式，请先另存为纯文字后再导入。')
        finally:workbook.close()
        repo=CopyRepository(self._engine)
        preview=classify_rows(repo,read_excel_rows(path))
        touched={d.id for group in (preview.new,preview.duplicate,preview.updates) for d in group}
        baseline={id:repo.get(id) for id in touched}
        token=secrets.token_hex(16)
        with self._lock:
            if len(self._previews)>=8:self._previews.pop(next(iter(self._previews)))
            self._previews[token]=(preview,baseline)
        return {'token':token,'new_count':len(preview.new),'duplicate_count':len(preview.duplicate),
                'update_count':len(preview.updates),'updates':[dict(id=d.id,title=d.title) for d in preview.updates],
                'errors':[{'row':e.row_number,'reason':e.reason} for e in preview.errors]}

    def apply(self,token,approved_ids,category_ids):
        if type(token) is not str or len(token)!=32 or type(approved_ids) is not list or type(category_ids) is not list or not all(type(x) is str for x in approved_ids+category_ids):
            raise ValueError('导入请求无效。')
        from yijianjifa.features.library.copy_service import single_category_ids
        category_ids = single_category_ids(category_ids)
        payload=json.dumps({'token':token,'approved':sorted(set(approved_ids)),'categories':sorted(set(category_ids))},ensure_ascii=False,sort_keys=True)
        digest=hashlib.sha256(payload.encode()).hexdigest();rid='excel_'+token
        with self._lock,self._engine.connect() as connection:
            connection.exec_driver_sql('BEGIN IMMEDIATE')
            try:
                row=connection.execute(text('SELECT payload_hash,response_json FROM web_command_receipts WHERE request_id=:id'),{'id':rid}).first()
                if row:
                    if row[0]!=digest:raise ValueError('这个预检查已经按其他选择完成，请重新检查表格。')
                    result=json.loads(row[1]);connection.commit();return result
                if token not in self._previews:raise ValueError('预检查已失效，请重新选择表格。')
                preview,baseline=self._previews[token]
                if not set(approved_ids)<={d.id for d in preview.updates}:raise ValueError('存在未预检查的更新。')
                categories=CategoryRepository(connection)
                if not set(category_ids)<={c.id for c in categories.list('copy')}:raise ValueError('分类已失效，请重新选择。')
                repo=CopyRepository(connection)
                if any(repo.get(id)!=before for id,before in baseline.items()):raise ValueError('资料在预检查后已修改，请重新检查表格。')
                result_core=apply_import(repo,preview,approved_update_ids=set(approved_ids))
                for id in set(result_core.new_ids+result_core.updated_ids+result_core.skipped_duplicate_ids):
                    if category_ids and repo.get(id) is not None:repo.set_categories(id,category_ids)
                result={'new_count':len(result_core.new_ids),'updated_count':len(result_core.updated_ids),
                        'duplicate_count':len(result_core.skipped_duplicate_ids),'error_count':result_core.error_count}
                connection.execute(text('INSERT INTO web_command_receipts(request_id,payload_hash,response_json) VALUES (:id,:hash,:response)'),{'id':rid,'hash':digest,'response':json.dumps(result)})
                connection.commit()
                return result
            except Exception:
                connection.rollback();raise
