from __future__ import annotations

from dataclasses import dataclass
import json
from pathlib import Path
import re
import shutil
import tempfile
import threading
import time
import uuid

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlalchemy import Engine

from yijianjifa.features.image_posts.service import IMAGE_SUFFIXES, ImagePostService, inspect_image, natural_key, sha256_file
from yijianjifa.os_integration.paths import AppPaths
from yijianjifa.storage.db import session_scope
from yijianjifa.storage.models import Setting

METADATA_COLUMNS = ("图片文件夹名称", "标题", "正文", "标签")
TOKEN_LIFETIME = 30 * 60
MAX_WORKBOOK_BYTES = 10 * 1024 * 1024
MAX_WORKBOOK_ROWS = 10_000
MAX_BATCH_GROUPS = 500


@dataclass
class _Stage:
    kind: str
    created: float
    values: list


def _selected(value) -> list[Path]:
    if value is None:
        return []
    if isinstance(value, (str, Path)):
        value = [value]
    return [Path(path).expanduser().absolute() for path in value]


def _safe_name(value: str) -> str:
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(value)).strip(" .")[:120]
    return name or "图文"


class ImagePostFiles:
    """Native-dialog inputs stay here; staging results expose names, never paths."""

    def __init__(self, engine: Engine, paths: AppPaths):
        self.service = ImagePostService(engine, paths)
        self._stages: dict[str, _Stage] = {}
        self._lock = threading.RLock()

    def _stage(self, kind, values):
        with self._lock:
            now = time.monotonic()
            self._stages = {token: stage for token, stage in self._stages.items() if now - stage.created < TOKEN_LIFETIME}
            if len(self._stages) >= 20:
                oldest = min(self._stages, key=lambda token: self._stages[token].created)
                self._stages.pop(oldest)
            token = str(uuid.uuid4())
            self._stages[token] = _Stage(kind, now, values)
            return token

    def _pending(self, token, kind):
        with self._lock:
            stage = self._stages.get(str(token))
        if stage is None or stage.kind != kind or time.monotonic() - stage.created >= TOKEN_LIFETIME:
            raise ValueError("导入预览已失效，请重新选择文件")
        return stage.values

    def cancel_stage(self, token):
        with self._lock:
            self._stages.pop(str(token), None)

    def stage_import(self, selected, batch=False):
        selected = _selected(selected)
        if not selected:
            return {"token": "", "groups": [], "errors": [], "cancelled": True}
        candidates, errors = [], []
        if batch:
            for root in selected:
                if root.is_symlink() or any(parent.is_symlink() for parent in root.parents) or not root.is_dir():
                    errors.append(f"{root.name}：请选择原始文件夹，不支持符号链接")
                    continue
                children = sorted(root.iterdir(), key=natural_key)
                for child in children:
                    if child.is_symlink():
                        errors.append(f"{child.name}：不支持符号链接")
                    elif child.is_dir():
                        files = sorted((path for path in child.iterdir() if path.suffix.lower() in IMAGE_SUFFIXES), key=natural_key)
                        if files:
                            candidates.append((child.name, child.name, files))
                        else:
                            errors.append(f"{child.name}：文件夹中没有支持的图片")
                root_files = [path for path in children if path.suffix.lower() in IMAGE_SUFFIXES and not path.is_dir()]
                if root_files:
                    candidates.append((root.name, f"{root.name}（根目录图片）", root_files))
        else:
            direct = []
            for selected_path in selected:
                if selected_path.is_dir() and not selected_path.is_symlink():
                    direct.extend(path for path in selected_path.iterdir() if path.suffix.lower() in IMAGE_SUFFIXES)
                else:
                    direct.append(selected_path)
            direct = sorted(direct, key=natural_key)
            if direct:
                name = selected[0].name if len(selected) == 1 and selected[0].is_dir() else direct[0].stem
                candidates.append((name, name, direct))
        if len(candidates) > MAX_BATCH_GROUPS:
            return {"token": "", "groups": [], "errors": [f"一次最多导入 {MAX_BATCH_GROUPS} 篇图文，请分批选择"], "cancelled": False}
        values, groups = [], []
        for key, name, paths in candidates:
            try:
                if not 1 <= len(paths) <= 100:
                    raise ValueError("一篇图文需要 1 至 100 张图片")
                checked = []
                for path in paths:
                    source, _image = inspect_image(path)
                    _image.close()
                    checked.append((source, sha256_file(source)))
                group = {"key": key, "name": name, "image_count": len(paths), "filenames": [path.name for path in paths], "title": "", "body": "", "tags": ""}
                values.append({"group": group, "sources": checked})
                groups.append(group.copy())
            except (ValueError, OSError) as error:
                errors.append(f"{name}：{self._error(error)}")
        token = self._stage("import", values) if values else ""
        return {"token": token, "groups": groups, "errors": errors, "cancelled": False}

    def confirm_import(self, token, *, receipt_writer=None):
        return self._confirm_atomic(token, "import", self._import_pending, receipt_writer)

    def _confirm_atomic(self, token, kind, operation, receipt_writer):
        """Business rows, token recovery, and bridge receipt share one commit.

        Failed confirmation leaves the private stage available for retry. A
        committed token has a path-free durable result even after app restart.
        Files made by a rejected transaction are removed after rollback.
        """
        token = str(token)
        key = f"image_posts.confirm:{kind}:{token}"
        created_ids = []
        with self._lock:
            try:
                with session_scope(self.service.engine) as db:
                    previous = db.get(Setting, key)
                    if previous is not None:
                        result = json.loads(previous.value_json)
                    else:
                        try:
                            pending = self._pending(token, kind)
                        except ValueError as error:
                            result = {"imported" if kind == "import" else "updated": 0, "ids": [], "errors": [str(error)], "cancelled": False}
                        else:
                            result = operation(db, pending, created_ids)
                        db.add(Setting(key=key, value_json=json.dumps(result, ensure_ascii=False)))
                    if receipt_writer is not None:
                        receipt_writer(db, result)
                    db.flush()
            except Exception:
                for post_id in created_ids:
                    self.service._remove_directory(post_id)
                raise
            # Consumption follows commit, so receipt failure cannot lose grants.
            self._stages.pop(token, None)
            return result

    def _import_pending(self, db, values, created_ids):
        ids, errors = [], []
        for value in values:
            group = value["group"]
            try:
                for source, expected in value["sources"]:
                    # Refuse stale previews and links introduced after selection.
                    _source, _image = inspect_image(source)
                    _image.close()
                    if sha256_file(source) != expected:
                        raise ValueError(f"{source.name}：预览后图片发生变化，请重新选择")
                post = self.service.import_group([source for source, _ in value["sources"]], display_name=group["name"], title=group["title"], body=group["body"], tags=group["tags"], _session=db)
                ids.append(post.id)
                created_ids.append(post.id)
            except (ValueError, OSError) as error:
                errors.append(f"{group['name']}：{self._error(error)}")
        return {"imported": len(ids), "ids": ids, "errors": errors, "cancelled": False}

    def stage_metadata(self, selected):
        selected = _selected(selected)
        if not selected:
            return {"token": "", "matches": [], "errors": [], "cancelled": True}
        if len(selected) != 1:
            return {"token": "", "matches": [], "errors": ["请选择一份图文文案 Excel"], "cancelled": False}
        path = selected[0]
        matches, errors = [], []
        try:
            if path.is_symlink() or any(parent.is_symlink() for parent in path.parents) or not path.is_file() or path.suffix.lower() != ".xlsx":
                raise ValueError("请选择原始 .xlsx 文件，不支持符号链接")
            if path.stat().st_size > MAX_WORKBOOK_BYTES:
                raise ValueError("Excel 大小超过 10 MB")
            workbook = load_workbook(path, read_only=True, data_only=False)
            try:
                sheet = workbook.active
                if sheet.max_row > MAX_WORKBOOK_ROWS:
                    raise ValueError("Excel 行数超过 10000 行，请分批导入")
                rows = sheet.iter_rows()
                header = next(rows, ())
                if tuple(str(cell.value or "").strip() for cell in header[:4]) != METADATA_COLUMNS:
                    raise ValueError("Excel 表头不正确，请使用图文文案模板")
                catalog = {}
                offset = 0
                while True:
                    records, total = self.service.list(offset=offset, limit=500)
                    for record in records:
                        catalog.setdefault(record.display_name, []).append(record.id)
                    offset += len(records)
                    if offset >= total or not records:
                        break
                entries = []
                for row_number, cells in enumerate(rows, start=2):
                    if not any(cell.value is not None for cell in cells):
                        continue
                    if any(cell.data_type == "f" for cell in cells):
                        errors.append(f"第 {row_number} 行：不支持公式，请粘贴实际文字")
                        continue
                    fields = [str(cell.value or "") for cell in cells[:4]]
                    fields.extend([""] * (4 - len(fields)))
                    name = fields[0].strip()
                    if not name:
                        errors.append(f"第 {row_number} 行：缺少图片文件夹名称")
                        continue
                    entries.append((row_number, name, fields[1:]))
                duplicates = {name for _, name, _ in entries if sum(1 for _, value, _ in entries if value == name) > 1}
                for row_number, name, fields in entries:
                    ids = catalog.get(name, [])
                    if name in duplicates:
                        errors.append(f"第 {row_number} 行：{name} 在 Excel 中重复，请保留一行")
                    elif not ids:
                        errors.append(f"第 {row_number} 行：{name} 未匹配到图文，请先导入对应文件夹")
                    elif len(ids) > 1:
                        errors.append(f"第 {row_number} 行：{name} 的图文名称重复，请先修改名称")
                    else:
                        matches.append({"id": ids[0], "key": name, "name": name, "title": fields[0], "body": fields[1], "tags": fields[2]})
            finally:
                workbook.close()
        except Exception as error:
            # Workbook parser errors are shown as a readable filename-only error.
            errors.append(f"{path.name}：{self._error(error)}")
        token = self._stage("metadata", matches) if matches else ""
        return {"token": token, "matches": matches, "errors": errors, "cancelled": False}

    def confirm_metadata(self, token, *, receipt_writer=None):
        return self._confirm_atomic(token, "metadata", self._metadata_pending, receipt_writer)

    def _metadata_pending(self, db, matches, _created_ids):
        ids, errors = [], []
        for match in matches:
            current = self.service.get(match["id"], _session=db)
            if current is None or current.display_name != match["name"]:
                errors.append(f"{match['name']}：图文已删除或改名，请重新导入文案")
                continue
            self.service.save(current.id, {field: match[field] for field in ("title", "body", "tags")}, _session=db)
            ids.append(current.id)
        return {"updated": len(ids), "ids": ids, "errors": errors, "cancelled": False}

    def export_template(self, selected):
        selected = _selected(selected)
        if not selected:
            return {"exported": 0, "errors": [], "cancelled": True}
        if len(selected) != 1:
            return {"exported": 0, "errors": ["请选择一个模板保存位置"], "cancelled": False}
        path = selected[0]
        if path.is_symlink() or any(parent.is_symlink() for parent in path.parents):
            return {"exported": 0, "errors": ["不支持符号链接，请选择原始保存位置"], "cancelled": False}
        if path.suffix.lower() != ".xlsx":
            path = path.with_suffix(".xlsx")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "图文文案"
        sheet.append(list(METADATA_COLUMNS))
        sheet.append(["产品A", "产品A的发布标题", "产品介绍正文", "#产品展示"])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column, width in (("A", 26), ("B", 36), ("C", 70), ("D", 36)):
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A2"
        try:
            workbook.save(path)
        except OSError as error:
            return {"exported": 0, "errors": [self._error(error)], "cancelled": False}
        finally:
            workbook.close()
        return {"exported": 1, "filename": path.name, "errors": [], "cancelled": False}

    def export(self, ids, selected):
        selected = _selected(selected)
        if not selected:
            return {"exported": 0, "errors": [], "cancelled": True}
        target = selected[0]
        if len(selected) != 1 or target.is_symlink() or any(parent.is_symlink() for parent in target.parents) or target.resolve().is_relative_to(self.service.root.resolve()):
            return {"exported": 0, "errors": ["请选择软件图片目录以外的原始文件夹"], "cancelled": False}
        target.mkdir(parents=True, exist_ok=True)
        exported, errors = 0, []
        for post_id in dict.fromkeys(ids):
            post = self.service.get(post_id)
            if post is None:
                errors.append("有一篇图文已删除，请刷新后重新选择")
                continue
            staging = Path(tempfile.mkdtemp(prefix=".image-post-export-", dir=target))
            try:
                names = []
                for image in post.images:
                    if not image.path.is_file() or sha256_file(image.path) != image.content_hash:
                        raise ValueError(f"{image.original_name}：软件图片副本缺失或已改变")
                    name = f"{image.position + 1:02d}{image.path.suffix.lower()}"
                    shutil.copy2(image.path, staging / name)
                    names.append(name)
                payload = {"display_name": post.display_name, "title": post.title, "body": post.body, "tags": post.tags, "images": names}
                (staging / "metadata.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
                destination = target / _safe_name(post.display_name)
                if destination.exists():
                    destination = target / f"{_safe_name(post.display_name)}-{post.id[:8]}"
                counter = 2
                while destination.exists():
                    destination = target / f"{_safe_name(post.display_name)}-{post.id[:8]}-{counter}"
                    counter += 1
                staging.rename(destination)
                exported += 1
            except (ValueError, OSError) as error:
                errors.append(f"{post.display_name}：{self._error(error)}")
                shutil.rmtree(staging, ignore_errors=True)
        return {"exported": exported, "errors": errors, "cancelled": False}

    @staticmethod
    def _error(error):
        if isinstance(error, ValueError):
            return str(error)
        if isinstance(error, OSError):
            return "文件无法读取或保存，请检查文件权限和磁盘空间"
        return "文件内容无法读取，请重新保存为标准 Excel 文件"
