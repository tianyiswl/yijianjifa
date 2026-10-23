from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
import uuid
from typing import Iterable, Mapping

from openpyxl import Workbook, load_workbook

from yijianjifa.features.library.copy_service import CopyDraft, CopyRepository

EXPECTED_COLUMNS = ["文案编号（ID）", "标题", "口播稿", "发布文案", "标签", "全程字幕"]


class _ExcelRow(dict[str, str]):
    """Keep the worksheet position without adding a seventh content field."""

    def __init__(self, values: dict[str, str], row_number: int):
        super().__init__(values)
        self.row_number = row_number


@dataclass(frozen=True, slots=True)
class ImportErrorRow:
    row_number: int
    reason: str
    raw: dict[str, str]


@dataclass(slots=True)
class ImportPreview:
    new: list[CopyDraft] = field(default_factory=list)
    duplicate: list[CopyDraft] = field(default_factory=list)
    updates: list[CopyDraft] = field(default_factory=list)
    errors: list[ImportErrorRow] = field(default_factory=list)


@dataclass(frozen=True, slots=True)
class ImportApplyResult:
    new_ids: list[str]
    updated_ids: list[str]
    skipped_duplicate_ids: list[str]
    error_count: int


def _normalize_row(row: Mapping[str, object]) -> dict[str, str]:
    return {column: str(row.get(column, "") or "").strip() for column in EXPECTED_COLUMNS}


def _draft_from_row(row: Mapping[str, object], *, generated_id: str | None = None) -> CopyDraft:
    normalized = _normalize_row(row)
    return CopyDraft(
        id=normalized[EXPECTED_COLUMNS[0]] or generated_id or _new_id(),
        title=normalized[EXPECTED_COLUMNS[1]],
        voice=normalized[EXPECTED_COLUMNS[2]],
        body=normalized[EXPECTED_COLUMNS[3]],
        tags=normalized[EXPECTED_COLUMNS[4]],
        full=normalized[EXPECTED_COLUMNS[5]],
    )


def _new_id() -> str:
    return "W" + uuid.uuid4().hex[:12].upper()


def classify_rows(copy_repo: CopyRepository, rows: Iterable[Mapping[str, object]]) -> ImportPreview:
    raw_rows = [
        (row.row_number if isinstance(row, _ExcelRow) else index, _normalize_row(row))
        for index, row in enumerate(rows, start=2)
    ]
    preview = ImportPreview()

    # Same explicit ID pointing to different content in one import is a blocking row error.
    by_id: dict[str, list[tuple[int, dict[str, str]]]] = {}
    for index, row in raw_rows:
        copy_id = row[EXPECTED_COLUMNS[0]]
        if copy_id:
            by_id.setdefault(copy_id, []).append((index, row))

    conflicting_rows: set[int] = set()
    for copy_id, group in by_id.items():
        content_variants = {
            tuple(row[col] for col in EXPECTED_COLUMNS[1:])
            for _, row in group
        }
        if len(content_variants) > 1:
            for row_number, raw in group:
                conflicting_rows.add(row_number)
                preview.errors.append(
                    ImportErrorRow(
                        row_number=row_number,
                        reason=f"同一文案编号 {copy_id} 在本次表格中指向不同内容",
                        raw=raw,
                    )
                )

    seen_exact: set[tuple[str, str, str, str, str, str]] = set()
    for row_number, row in raw_rows:
        if row_number in conflicting_rows:
            continue
        exact_key = tuple(row[col] for col in EXPECTED_COLUMNS)
        if exact_key in seen_exact:
            preview.duplicate.append(_draft_from_row(row))
            continue
        seen_exact.add(exact_key)

        explicit_id = row[EXPECTED_COLUMNS[0]]
        if explicit_id:
            draft = _draft_from_row(row)
            existing = copy_repo.get(explicit_id)
            if existing is None:
                preview.new.append(draft)
            elif existing.content_tuple == draft.content_tuple:
                preview.duplicate.append(existing)
            else:
                preview.updates.append(draft)
            continue

        draft = _draft_from_row(row)
        duplicate = copy_repo.find_by_content(draft)
        if duplicate is not None:
            preview.duplicate.append(duplicate)
        else:
            preview.new.append(draft)

    return preview


def apply_import(
    copy_repo: CopyRepository,
    preview: ImportPreview,
    *,
    approved_update_ids: set[str],
) -> ImportApplyResult:
    new_ids: list[str] = []
    updated_ids: list[str] = []
    for draft in preview.new:
        copy_repo.create(
            id=draft.id,
            title=draft.title,
            voice=draft.voice,
            body=draft.body,
            tags=draft.tags,
            full=draft.full,
        )
        new_ids.append(draft.id)
    for draft in preview.updates:
        if draft.id not in approved_update_ids:
            continue
        copy_repo.update(draft)
        updated_ids.append(draft.id)
    return ImportApplyResult(
        new_ids=new_ids,
        updated_ids=updated_ids,
        skipped_duplicate_ids=[draft.id for draft in preview.duplicate],
        error_count=len(preview.errors),
    )


def read_excel_rows(path: str | Path, *, sheet_name: str | None = None) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        return []
    headers = [str(value or "").strip() for value in values[0]]
    if headers[: len(EXPECTED_COLUMNS)] != EXPECTED_COLUMNS:
        raise ValueError("Excel 表头必须严格使用六列文案模板")
    rows: list[dict[str, str]] = []
    for row_number, values_row in enumerate(values[1:], start=2):
        if not any(value not in (None, "") for value in values_row[: len(EXPECTED_COLUMNS)]):
            continue
        rows.append(
            _ExcelRow(
                {
                    column: str(values_row[index] or "").strip()
                    for index, column in enumerate(EXPECTED_COLUMNS)
                },
                row_number,
            )
        )
    return rows


def export_template(path: str | Path, *, include_example: bool = True) -> Path:
    target = Path(path)
    workbook = Workbook()
    template = workbook.active
    template.title = "文案模板"
    template.append(EXPECTED_COLUMNS)
    sample = ["", "先看看整体", "先从整体看起，再把镜头拉近。远处看布局，近处看细节。", "先看整体，再看细节。你更喜欢哪种镜头？", "日常记录 镜头观察", "换个角度，发现细节"]
    template.append(sample if include_example else ["", "", "", "", "", ""])

    if include_example:
        example = workbook.create_sheet("填写示例")
        example.append(EXPECTED_COLUMNS)
        example.append(["", "便携水杯", "今天介绍这款水杯。", "轻巧便携，日常使用方便。", "水杯 便携", "新品推荐"])

    instructions = workbook.create_sheet("填写说明")
    instructions.append(["字段", "说明"])
    instructions.append(["填写示例", "首页第二行是示例，请替换为自己的文案，或删除示例后填写；保留示例会将其一并导入。"])
    instructions.append(["文案编号（ID）", "首次导入可留空，由程序生成；后续导出带编号表可用于更新。"])
    instructions.append(["口播稿", "可留空，表示本条制作时不配音。"])
    instructions.append(["全程字幕", "可留空；分类不写入 Excel，在程序中设置。"])
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    return target


def export_with_ids(copy_repo: CopyRepository, copy_ids: Iterable[str], path: str | Path) -> Path:
    target = Path(path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "文案"
    sheet.append(EXPECTED_COLUMNS)
    for draft in copy_repo.list_by_ids(copy_ids):
        sheet.append([draft.id, draft.title, draft.voice, draft.body, draft.tags, draft.full])
        # User copy is literal text, including leading equals signs.
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    return target
