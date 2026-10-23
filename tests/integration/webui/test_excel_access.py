import json
import pytest
import yijianjifa.app_services as module
from yijianjifa.features.library.copy_service import CopyRepository,CopyDraft,CategoryRepository
from yijianjifa.features.library.excel_import import EXPECTED_COLUMNS
from openpyxl import Workbook


def access(engine):
    cls=getattr(module,'ExcelAccess',None)
    assert cls is not None,'Excel preview capability service is missing'
    return cls(engine)


def book(tmp_path,rows):
    path=tmp_path/'文案.xlsx';w=Workbook();w.active.append(EXPECTED_COLUMNS)
    for row in rows:w.active.append(row)
    w.save(path);return path


def test_preview_no_write_apply_only_approved_updates_and_retry(db_engine,tmp_path):
    repo=CopyRepository(db_engine);repo.create(id='old',title='original',voice='',body='',tags='',full='')
    api=access(db_engine);p=api.preview(book(tmp_path,[['old','changed','','','',''],['','new','','','','']]))
    assert len(repo.list_all())==1 and repo.get('old').title=='original'
    assert p['new_count']==1 and p['update_count']==1
    result=api.apply(p['token'],[],[])
    assert result['new_count']==1 and result['updated_count']==0
    assert api.apply(p['token'],[],[])==result
    assert len(repo.list_all())==2 and repo.get('old').title=='original'
    with pytest.raises(ValueError):api.apply(p['token'],['old'],[])


def test_stale_preview_cannot_overwrite_concurrent_change(db_engine,tmp_path):
    repo=CopyRepository(db_engine);repo.create(id='old',title='before',voice='',body='',tags='',full='')
    api=access(db_engine);p=api.preview(book(tmp_path,[['old','from sheet','','','','']]))
    repo.update(CopyDraft('old','new edit','','','',''))
    with pytest.raises(ValueError):api.apply(p['token'],['old'],[])
    assert repo.get('old').title=='new edit'


def test_categories_preserved_and_added_for_duplicate(db_engine,tmp_path):
    repo=CopyRepository(db_engine);repo.create(id='old',title='same',voice='',body='',tags='',full='')
    cats=CategoryRepository(db_engine);a=cats.create('copy','A');b=cats.create('copy','B');repo.set_categories('old',[a.id])
    api=access(db_engine);p=api.preview(book(tmp_path,[['old','same','','','','']]))
    api.apply(p['token'],[],[b.id])
    assert repo.category_ids('old')=={a.id,b.id}


def test_preview_token_not_path_and_invalid_approval_denied(db_engine,tmp_path):
    api=access(db_engine);path=book(tmp_path,[['','new','','','','']]);p=api.preview(path)
    assert str(tmp_path) not in json.dumps(p)
    with pytest.raises(ValueError):api.apply(str(path),[],[])
    with pytest.raises(ValueError):api.apply(p['token'],['not-in-preview'],[])
    assert CopyRepository(db_engine).list_all()==[]


def test_formula_workbook_rejected_without_evaluation(db_engine,tmp_path):
    api=access(db_engine);path=book(tmp_path,[['','=1+1','','','','']])
    with pytest.raises(ValueError):api.preview(path)
    assert CopyRepository(db_engine).list_all()==[]


def test_export_user_text_is_not_an_excel_formula(db_engine,tmp_path):
    from openpyxl import load_workbook
    from yijianjifa.features.library.excel_import import export_with_ids
    repo=CopyRepository(db_engine);repo.create(id='C1',title='=1+1',voice='',body='',tags='',full='')
    path=export_with_ids(repo,['C1'],tmp_path/'out.xlsx')
    workbook=load_workbook(path,data_only=False)
    try:
        assert workbook.active['B2'].value=='=1+1'
        assert workbook.active['B2'].data_type=='s'
    finally:workbook.close()


@pytest.mark.parametrize('blank_position,expected_error_rows', [(0, [3, 4]), (1, [2, 4])])
def test_preview_preserves_excel_error_row_numbers_across_blank_rows(
    db_engine, tmp_path, blank_position, expected_error_rows
):
    rows = [
        ['conflict', 'first content', '', '', '', ''],
        ['conflict', 'second content', '', '', '', ''],
        ['', 'valid new', '', '', '', ''],
    ]
    rows.insert(blank_position, ['', '', '', '', '', ''])
    path = book(tmp_path, rows)
    api = access(db_engine)

    preview = api.preview(path)

    assert [error['row'] for error in preview['errors']] == expected_error_rows
    assert preview['new_count'] == 1
    assert preview['duplicate_count'] == preview['update_count'] == 0
    result = api.apply(preview['token'], [], [])
    assert result == {'new_count': 1, 'updated_count': 0, 'duplicate_count': 0, 'error_count': 2}
    assert [draft.title for draft in CopyRepository(db_engine).list_all()] == ['valid new']
