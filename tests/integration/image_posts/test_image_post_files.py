import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from PySide6.QtGui import QImage

from yijianjifa.features.image_posts.files import ImagePostFiles, METADATA_COLUMNS


def picture(path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    image = QImage(20, 30, QImage.Format.Format_RGB32)
    image.fill(0xFF005588)
    assert image.save(str(path))
    return path


def test_batch_staging_no_paths_root_separate_and_natural_names(db_engine, app_paths, tmp_path):
    root = tmp_path / "batch"
    for name in ("产品10", "产品2", "产品1"):
        for number in (10, 2, 1):
            picture(root / name / f"{number}.png")
    picture(root / "散图.png")
    files = ImagePostFiles(db_engine, app_paths)
    staged = files.stage_import([root], batch=True)
    assert [group["name"] for group in staged["groups"]] == ["产品1", "产品2", "产品10", "batch（根目录图片）"]
    assert staged["groups"][0]["filenames"] == ["1.png", "2.png", "10.png"]
    assert str(tmp_path) not in json.dumps(staged, ensure_ascii=False)
    assert files.service.list()[1] == 0
    confirmed = files.confirm_import(staged["token"])
    assert confirmed["imported"] == 4 and len(confirmed["ids"]) == 4
    assert files.confirm_import(staged["token"]) == confirmed
    assert files.service.list()[1] == 4


def test_metadata_matches_names_not_rows_and_rejects_ambiguous(db_engine, app_paths, tmp_path):
    files = ImagePostFiles(db_engine, app_paths)
    first = files.service.import_group([picture(tmp_path / "1.png")], display_name="产品A")
    second = files.service.import_group([picture(tmp_path / "2.png")], display_name="产品B")
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(METADATA_COLUMNS))
    sheet.append(["产品B", "B标题", "B正文", "#B"])
    sheet.append(["不存在", "错位标题", "错位正文", ""])
    sheet.append(["产品A", "A标题", "A正文", "#A"])
    path = tmp_path / "metadata.xlsx"
    workbook.save(path)
    staged = files.stage_metadata([path])
    assert {match["id"] for match in staged["matches"]} == {first.id, second.id}
    assert any("不存在" in error for error in staged["errors"])
    assert str(tmp_path) not in json.dumps(staged, ensure_ascii=False)
    result = files.confirm_metadata(staged["token"])
    assert result["updated"] == 2
    assert files.service.get(first.id).title == "A标题"
    assert files.service.get(second.id).title == "B标题"
    files.service.import_group([picture(tmp_path / "3.png")], display_name="产品A")
    ambiguous = files.stage_metadata([path])
    assert len(ambiguous["matches"]) == 1
    assert any("重复" in error and "产品A" in error for error in ambiguous["errors"])


def test_changed_import_source_refused_and_cancellation(db_engine, app_paths, tmp_path):
    files = ImagePostFiles(db_engine, app_paths)
    source = picture(tmp_path / "1.png")
    staged = files.stage_import([source])
    source.write_bytes(b"changed after staging")
    result = files.confirm_import(staged["token"])
    assert result["imported"] == 0 and result["errors"]
    assert files.stage_import([])["cancelled"]
    assert files.stage_metadata([])["cancelled"]


def test_template_and_export_include_ordered_images_and_metadata(db_engine, app_paths, tmp_path):
    files = ImagePostFiles(db_engine, app_paths)
    template = tmp_path / "模板.xlsx"
    result = files.export_template([template])
    assert result["exported"] == 1
    assert tuple(next(load_workbook(template, read_only=True).active.values)) == METADATA_COLUMNS
    post = files.service.import_group([picture(tmp_path / "2.png"), picture(tmp_path / "1.png")], display_name="产品/名称", title="=标题", body="正文", tags="#标签")
    target = tmp_path / "export"
    result = files.export([post.id], [target])
    assert result["exported"] == 1
    metadata = list(target.rglob("metadata.json"))
    assert len(metadata) == 1
    payload = json.loads(metadata[0].read_text(encoding="utf-8"))
    assert payload["title"] == "=标题"
    assert payload["images"] == ["01.jpg", "02.jpg"]
    assert all((metadata[0].parent / name).exists() for name in payload["images"])
