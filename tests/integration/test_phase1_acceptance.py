from pathlib import Path

from openpyxl import load_workbook

from yijianjifa.features.library.copy_service import CopyRepository
from yijianjifa.features.library.excel_import import EXPECTED_COLUMNS, export_template
from yijianjifa.features.library.media_storage import MediaStorageService
from yijianjifa.features.library.usage_progress import UsageProgressService


def test_rename_delete_reimport_keeps_historical_identity(db_engine, app_paths, tmp_path):
    media = MediaStorageService(db_engine, app_paths, kind="video")
    source = tmp_path / "a.mp4"
    source.write_bytes(b"same")
    first = media.import_media(source, keep_copy=False, category_ids=[])
    media.rename(first.id, "展示名称")
    media.delete_media_record(first.id)

    renamed = tmp_path / "renamed.mp4"
    renamed.write_bytes(b"same")
    second = media.import_media(renamed, keep_copy=False, category_ids=[])

    assert second.display_name == "renamed.mp4"
    assert media.history.identity_seen(second.content_hash)


def test_excel_template_has_exact_six_columns_and_no_category_column(tmp_path):
    target = export_template(tmp_path / "copy-template.xlsx")
    wb = load_workbook(target, read_only=True)
    headers = [cell.value for cell in next(wb["文案模板"].iter_rows())]
    assert headers == EXPECTED_COLUMNS
    assert all("分类" not in header for header in headers)


def test_copy_updates_do_not_merge_usage_progress_scopes(db_engine):
    copies = CopyRepository(db_engine)
    copies.create(id="W001", title="标题", voice="", body="发布文案", tags="标签", full="")
    progress = UsageProgressService(db_engine)
    progress.mark_used("mixing", "W001")
    assert progress.is_used("mixing", "W001")
    assert not progress.is_used("finished_metadata", "W001")
